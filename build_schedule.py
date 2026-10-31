import datetime
import hashlib
import json
import os
import re
import subprocess
from docx import Document
import requests

COURSES_CONFIG = {
    "1": {
        "url": "https://www.dropbox.com/scl/fo/ltk4mvd0monarv8c9k70a/ACkK0WsKu-WA81W7FYHt3YA/1kurs.doc?rlkey=yp9pp95tget8jo9oct0yppp7n&st=971tfg0i&dl=0",
        "raw_file": "1kurs.doc",
        "docx_file": "1kurs.docx",
        "needs_conversion": True,
    },
    "2": {
        "url": "https://www.dropbox.com/scl/fo/ltk4mvd0monarv8c9k70a/ANeggPvlbzFonxQ6QxDE0qQ/2kurs.doc?rlkey=yp9pp95tget8jo9oct0yppp7n&st=funb9d22&dl=0",
        "raw_file": "2kurs.doc",
        "docx_file": "2kurs.docx",
        "needs_conversion": True,
    },
    "3": {
        "url": "https://www.dropbox.com/scl/fo/ltk4mvd0monarv8c9k70a/AOtlH0FGL7cybx5XbUDT2IM/3kurs.doc?rlkey=yp9pp95tget8jo9oct0yppp7n&st=s92322wo&dl=0",
        "raw_file": "3kurs.doc",
        "docx_file": "3kurs.docx",
        "needs_conversion": True,
    },
    "4": {
        "url": "https://www.dropbox.com/scl/fo/ltk4mvd0monarv8c9k70a/AA480GB96NvLPj1FCveHNO4/4kurs.docx?rlkey=yp9pp95tget8jo9oct0yppp7n&st=g9bhzpdc&dl=0",
        "raw_file": "4kurs.docx",
        "docx_file": "4kurs.docx",
        "needs_conversion": False,
    },
}

def format_datetime_ru(dt):
    months = ["янв", "фев", "мар", "апр", "мая", "июн", "июл", "авг", "сен", "окт", "ноя", "дек"]
    return f"{dt.day} {months[dt.month - 1]} в {dt.strftime('%H:%M')}"

def download_file(course_id, cfg):
    url = cfg["url"].replace("&dl=0", "&dl=1").replace("?dl=0", "?dl=1")
    cache_buster = int(datetime.datetime.now().timestamp())
    url = f"{url}&_ts={cache_buster}"
    headers = {
        "User-Agent": "Mozilla/5.0",
        "Cache-Control": "no-cache, no-store, must-revalidate",
        "Pragma": "no-cache"
    }
    res = requests.get(url, headers=headers, timeout=30)
    res.raise_for_status()
    with open(cfg["raw_file"], "wb") as f:
        f.write(res.content)

def extract_decanat_modtime(docx_path):
    try:
        import docx
        doc = docx.Document(docx_path)
        dt = doc.core_properties.modified
        if dt:
            if dt.tzinfo is not None:
                minsk_tz = datetime.timezone(datetime.timedelta(hours=3))
                dt = dt.astimezone(minsk_tz)
            else:
                dt = dt + datetime.timedelta(hours=3)
            return format_datetime_ru(dt)
    except Exception as e:
        print(f"Не удалось извлечь время деканата из {docx_path}: {e}")
    return None

def convert_docs(files):
    if not files:
        return
    paths = ["/Applications/LibreOffice.app/Contents/MacOS/soffice", "soffice"]
    cmd = next((p for p in paths if os.path.exists(p) or subprocess.run(["which", p], stdout=subprocess.PIPE).returncode == 0), None)
    if not cmd:
        print("LibreOffice не найден")
        return
    subprocess.run([cmd, "--headless", "--convert-to", "docx", *files, "--outdir", "."], check=True)

def format_time(t):
    digits = re.sub(r"\s+", "", t)
    m = re.match(r"(\d{2})[:.]?(\d{2})[-–](\d{2})[:.]?(\d{2})", digits)
    return f"{m.group(1)}:{m.group(2)}–{m.group(3)}:{m.group(4)}" if m else digits

def is_cell_valid_for_slot(cell_text, cur_time):
    if not cell_text or not cur_time:
        return True
    # Strip calendar date ranges (e.g. 05.10-09.11) so they don't get misparsed as time of day
    cleaned = re.sub(r"\b\d{1,2}\.(?:0[1-9]|1[0-2])\s*[-–—]\s*\d{1,2}\.(?:0[1-9]|1[0-2])\b", "", cell_text)
    m = re.search(r"(\d{1,2})[.:](\d{2})\s*[-–—]\s*(\d{1,2})[.:](\d{2})", cleaned)
    if not m:
        m = re.search(r"\b(\d{2})(\d{2})\s*[-–—]\s*(\d{2})(\d{2})\b", cleaned)
    if not m:
        return True
    sh, sm = int(m.group(1)), int(m.group(2))
    eh, em = int(m.group(3)), int(m.group(4))
    start_min = sh * 60 + sm
    end_min = eh * 60 + em

    cur_m = re.match(r"(\d{1,2})[:.](\d{2})[–-](\d{1,2})[:.](\d{2})", cur_time)
    if not cur_m:
        return True
    c_sh, c_sm = int(cur_m.group(1)), int(cur_m.group(2))
    c_eh, c_em = int(cur_m.group(3)), int(cur_m.group(4))
    c_start_min = c_sh * 60 + c_sm
    c_end_min = c_eh * 60 + c_em

    # If cell ends before or at current slot start
    if end_min <= c_start_min + 15:
        return False
    # If cell starts after current slot ends
    if start_min >= c_end_min - 15:
        return False
    # If cell ends by 14:50 and cur_slot is 14:20-15:40 (5th pair slot)
    if end_min <= 14 * 60 + 50 and c_start_min >= 14 * 60 + 15:
        return False

    return True

def extract_pe_time(raw_text, cur_time=None):
    if not raw_text or not re.search(r"физическ.*культ|спорт\s*комплекс", raw_text, re.I):
        return None
    # 1. Ищем явное время в тексте ячейки для спорткомплекса на Октябрьской (основная группа)
    m_okt = re.search(r"(\d{1,2})[.:](\d{2})\s*[-–—]\s*(\d{1,2})[.:](\d{2})[^\n]*октябрьск", raw_text, re.I)
    if m_okt:
        return f"{int(m_okt.group(1)):02d}:{m_okt.group(2)}–{int(m_okt.group(3)):02d}:{m_okt.group(4)}"
    m_any = re.search(r"(\d{1,2})[.:](\d{2})\s*[-–—]\s*(\d{1,2})[.:](\d{2})", raw_text)
    if m_any:
        return f"{int(m_any.group(1)):02d}:{m_any.group(2)}–{int(m_any.group(3)):02d}:{m_any.group(4)}"
    # 2. Стандартные официальные сдвиги расписания спорткомплекса геофака БГУ
    if cur_time and ("14:20" in cur_time or "14:10" in cur_time):
        return "14:40–16:00"
    if cur_time and "12:50" in cur_time:
        return "13:10–14:30"
    if cur_time and "11:20" in cur_time:
        return "10:50–12:10"
    return None

def extract_custom_cell_time(raw_text, cur_time=None):
    if not raw_text:
        return None

    # 1. Physical education / Sports complex shifts
    if re.search(r"физическ.*культ|спорт\s*комплекс", raw_text, re.I):
        pe_t = extract_pe_time(raw_text, cur_time)
        if pe_t:
            return pe_t

    # 2. Chemistry faculty lab (e.g. Course 1 Group 8 Monday 11:50–14:50)
    if re.search(r"хим\.?\s*фак", raw_text, re.I) and re.search(r"общая\s+химия|лаб", raw_text, re.I):
        return "11:50–14:50"

    # 3. Strip calendar date ranges (e.g. 05.10-09.11)
    cleaned = re.sub(r"\b\d{1,2}\.(?:0[1-9]|1[0-2])\s*[-–—]\s*\d{1,2}\.(?:0[1-9]|1[0-2])\b", "", raw_text)

    # 4. Check for explicit times like 11.50-14.50 or 14:20–17:20 or 14.25-17.15 or 15.00 – 16.20
    m = re.search(r"\b(\d{1,2})[.:](\d{2})\s*[-–—]\s*(\d{1,2})[.:](\d{2})\b", cleaned)
    if not m:
        # Check for 4-digit format like 1150-1450 or 1600-1720 or 1120-1240 or 1730-1850 or 1930-2050
        m = re.search(r"\b(0[89]|1[0-9]|2[01])([0-5]\d)\s*[-–—]\s*(0[89]|1[0-9]|2[02])([0-5]\d)\b", cleaned)

    if m:
        sh, sm = int(m.group(1)), int(m.group(2))
        eh, em = int(m.group(3)), int(m.group(4))
        start_min = sh * 60 + sm
        end_min = eh * 60 + em
        duration = end_min - start_min
        if 8 <= sh <= 21 and 8 <= eh <= 22 and 40 <= duration <= 300:
            return f"{sh:02d}:{sm:02d}–{eh:02d}:{em:02d}"

    return None

def clean_cell_text(text):
    if not text:
        return ""
    text = text.replace('\xad', '').replace('\u200b', '')
    text = re.sub(r"([а-яА-ЯёЁa-zA-Z]+)[-–—]\s*\n\s*([а-яА-ЯёЁa-zA-Z]+)", r"\1\2", text)
    text = re.sub(r"(ОБЩЕЕ\s+ЗЕМЛЕВЕДЕ)\s*\n\s*(НИЕ)", r"\1\2", text, flags=re.I)
    text = re.sub(r"(ОБЩЕЕ)\s*\n\s*(ЗЕМЛЕВЕДЕНИЕ)", r"\1 \2", text, flags=re.I)
    text = re.sub(r"(ИСТОРИЧЕ)\s*\n\s*(СКАЯ)", r"\1\2", text, flags=re.I)
    text = re.sub(r"(ЯЗЫК\s+ПРОФЕССИО)\s*\n\s*(НАЛЬНОГО)", r"\1\2", text, flags=re.I)
    text = re.sub(r"(ГРУНТОВЕ)\s*\n\s*(ДЕНИЯ)", r"\1\2", text, flags=re.I)

    legit = [
        'социально-эконом',
        'физико-географ',
        'экономико-географ',
        'научно-технич',
        'аппаратно-программ',
        'гис-картограф',
        'гис-технолог',
        'соц-эконом',
        'препод-стаж',
        'препод.-стаж',
        'творонович-севрук',
    ]

    def protect(m):
        full = m.group(0)
        lower = full.lower()
        for l in legit:
            if l in lower:
                return full.replace('-', '__HYPH__').replace('–', '__HYPH__').replace('—', '__HYPH__')
        return full

    text = re.sub(r'[а-яА-ЯёЁa-zA-Z]+[-–—][а-яА-ЯёЁa-zA-Z]+', protect, text)
    while re.search(r'([а-яА-ЯёЁa-zA-Z])[-–—]+([а-яА-ЯёЁa-zA-Z])', text):
        text = re.sub(r'([а-яА-ЯёЁa-zA-Z])[-–—]+([а-яА-ЯёЁa-zA-Z])', r'\1\2', text)
    text = text.replace('__HYPH__', '-')

    typos = [
        (re.compile(r'гибробиолог', re.I), 'гидробиолог'),
        (re.compile(r'автоматизи-?рованые', re.I), 'автоматизированные'),
        (re.compile(r'геоинформа-?цинное', re.I), 'геоинформационное'),
        (re.compile(r'иностраный\s+язык', re.I), 'иностранный язык'),
        (re.compile(r'генезис\s+месторожде-?ние\s+полезных', re.I), 'генезис месторождений полезных'),
    ]
    for p, r in typos:
        text = p.sub(r, text)

    lines = [re.sub(r"\s+", " ", l).strip() for l in text.split("\n") if l.strip()]
    return "\n".join([l for l in lines if not (len(re.sub(r'[:. –—-]', '', l)) in (7, 8) and re.sub(r'[:. –—-]', '', l).isdigit())])

def parse_docx(path):
    if not os.path.exists(path):
        return {}
    doc = Document(path)
    res = {}
    for table in doc.tables:
        if len(table.rows) < 3:
            continue

        # Extract table grid column widths for physical coordinate matching
        grid = [int(c.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}w'))
                for c in table._tbl.tblGrid.findall('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}gridCol')]
        x_coords = [0]
        for w in grid:
            x_coords.append(x_coords[-1] + w)

        header_groups = []  # list of (group_name, start_x, end_x)
        h_idx = 0
        for r_idx in range(min(3, len(table.rows))):
            row = table.rows[r_idx]
            c_idx = 0
            for tc in row._tr.findall('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tc'):
                gs = tc.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tcPr/{http://schemas.openxmlformats.org/wordprocessingml/2006/main}gridSpan')
                span = int(gs.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val')) if gs is not None else 1
                cell_obj = Document(path).tables[0].rows[0].cells[0]  # dummy
                txt = re.sub(r"\s+", " ", row.cells[c_idx].text).strip().upper()
                x1 = x_coords[c_idx]
                x2 = x_coords[min(c_idx + span, len(x_coords) - 1)]
                c_idx += span
                if "ГРУППА" in txt:
                    g_name = txt.split("\n")[0].strip()
                    header_groups.append((g_name, x1, x2))
            if header_groups:
                h_idx = r_idx
                break

        if not header_groups:
            continue

        # Group rows by (day, time)
        day_time_rows = []
        cur_day, cur_time = None, None
        for row in table.rows[h_idx + 1:]:
            cells = row.cells
            if len(cells) < 3:
                continue
            day_txt = re.sub(r"\s+", " ", cells[0].text).strip().upper()
            if day_txt and "ВРЕМЯ" not in day_txt and "ДНИ" not in day_txt and len(day_txt) > 2:
                cur_day = day_txt
            if not cur_day:
                continue
            time_txt = cells[1].text.strip()
            if time_txt and any(c.isdigit() for c in time_txt):
                cur_time = format_time(time_txt)
            if not cur_time:
                continue

            if not day_time_rows or day_time_rows[-1][0] != cur_day or day_time_rows[-1][1] != cur_time:
                day_time_rows.append((cur_day, cur_time, [row]))
            else:
                day_time_rows[-1][2].append(row)

        for cur_day, cur_time, rows in day_time_rows:
            num_rows = len(rows)
            for r_sub_idx, row in enumerate(rows):
                c_idx = 0
                for tc in row._tr.findall('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tc'):
                    gs = tc.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tcPr/{http://schemas.openxmlformats.org/wordprocessingml/2006/main}gridSpan')
                    span = int(gs.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val')) if gs is not None else 1
                    cell_x1 = x_coords[c_idx]
                    cell_x2 = x_coords[min(c_idx + span, len(x_coords) - 1)]
                    raw_cell = row.cells[c_idx].text
                    c_idx += span

                    # Find which group columns this cell overlaps
                    matching_groups = []
                    for g_name, gx1, gx2 in header_groups:
                        g_width = gx2 - gx1
                        overlap = max(0, min(cell_x2, gx2) - max(cell_x1, gx1))
                        if overlap >= 0.3 * g_width:
                            matching_groups.append(g_name)

                    if not matching_groups:
                        continue

                    custom_time = extract_custom_cell_time(raw_cell, cur_time)
                    if not custom_time and not is_cell_valid_for_slot(raw_cell, cur_time):
                        continue

                    lesson = clean_cell_text(raw_cell)
                    if "СПЕЦИАЛЬНОСТЬ" in lesson.upper() or "НАПРАВЛЕНИЕ" in lesson.upper() or not lesson:
                        continue

                    # Subgroup detection:
                    # In university schedule tables, when a slot has 2 rows:
                    # Row 0 = Subgroup 1, Row 1 = Subgroup 2.
                    # If the cell is vertically merged across rows, both subgroups attend together.
                    is_vertically_merged = False
                    if num_rows == 2:
                        other_r_idx = 1 if r_sub_idx == 0 else 0
                        other_cell_text = clean_cell_text(rows[other_r_idx].cells[c_idx - span].text)
                        if other_cell_text == lesson:
                            is_vertically_merged = True

                    subgroup_prefix = ""
                    if num_rows == 2 and not is_vertically_merged:
                        if not re.search(r'русск|немец|франц|англ|иностранн|культ', lesson, re.I):
                            subgroup_prefix = f"{r_sub_idx + 1} подгруппа\n"

                    tagged_lesson = (subgroup_prefix + lesson) if (subgroup_prefix and not lesson.startswith(subgroup_prefix)) else lesson

                    slot_time = custom_time if custom_time else cur_time

                    for g_name in matching_groups:
                        res.setdefault(g_name, {}).setdefault(cur_day, {}).setdefault(slot_time, [])
                        if tagged_lesson:
                            # Separate multiple subjects crammed into a single cell by Decanat
                            sub_lessons = re.split(r"\n(?=(?:Немецкий язык|Русский язык как иностранный|ИНОСТРАННЫЙ ЯЗЫК|Французский язык))", tagged_lesson, flags=re.I)
                            for sub in sub_lessons:
                                sub = sub.strip()
                                # Ignore phantom entries consisting only of subgroup prefix or time
                                clean_sub = re.sub(r"^[12]\s*подгруппа\s*", "", sub, flags=re.I).strip()
                                clean_sub = re.sub(r"^\d{1,2}[.:]\d{2}\s*[-–—]\s*\d{1,2}[.:]\d{2}\s*", "", clean_sub).strip()
                                clean_sub = re.sub(r"^\d{4}\s*[-–—]\s*\d{4}\s*", "", clean_sub).strip()
                                if not clean_sub:
                                    continue
                                if sub and sub not in res[g_name][cur_day][slot_time]:
                                    res[g_name][cur_day][slot_time].append(sub)
    return res

def detect_lesson_language(text):
    if not text:
        return None
    lower = text.lower()
    # 1. French: teacher Гулезова or word франц
    if "гулезов" in lower or "франц" in lower:
        return "FR"
    # 2. German: teachers Пригодич, Хадасевич or word нем (avoiding неметаллическ)
    if ("нем" in lower and "неметал" not in lower) or "пригодич" in lower or "хадасевич" in lower:
        return "DE"
    # 3. Russian as foreign language (RKI)
    if "русск" in lower or "рки" in lower:
        if any(w in lower for w in ["иностр", "язык", "рки"]):
            return "RKI"
    # 4. English: teachers or explicit tokens or generic foreign language
    if any(t in lower for t in ["шарейко", "селиванова", "дворкина", "боровикова", "давыдова", "пышняк", "англ", "eduengscience"]):
        return "EN"
    if "иностранн" in lower or "язык профессионального общения" in lower or "практика иностранного языка" in lower:
        return "EN"
    return None

def propagate_stream_languages(courses_data):
    """
    In university schedules, foreign language stream classes (French, German)
    are often entered by Decanat in only 1 or 2 group columns, but attendable
    by students from any group of the same course who study that language.
    This function shares stream language lessons across groups of the same course
    if the group's slot is empty or only contains other foreign language / PE classes.
    """
    for cid, groups in courses_data.items():
        if not isinstance(groups, dict):
            continue
        stream_map = {}
        for g_name, days in groups.items():
            if not isinstance(days, dict):
                continue
            for day_key, slots in days.items():
                if not isinstance(slots, dict):
                    continue
                for time_str, lessons in slots.items():
                    for l in lessons:
                        lang = detect_lesson_language(l)
                        if lang in ("FR", "DE"):
                            key = (day_key, time_str, lang)
                            if key not in stream_map:
                                stream_map[key] = l

        for (day_key, time_str, lang), stream_lesson in stream_map.items():
            for g_name, days in groups.items():
                if not isinstance(days, dict) or day_key not in days:
                    continue
                # Group 9 Course 1 studies English at ул. Курчатова, 5 on Mon/Thu and has no classes at ChemFac/GeoFac
                if cid == "1" and g_name == "ГРУППА 9" and any(k in day_key.upper() for k in ["ПОНЕДЕЛЬНИК", "ЧЕТВЕРГ"]):
                    continue
                slots = days[day_key]
                current_lessons = slots.get(time_str, [])
                has_lang = any(detect_lesson_language(l) == lang for l in current_lessons)
                if has_lang:
                    continue
                is_available = not current_lessons or all(
                    detect_lesson_language(l) is not None or bool(re.search(r"физическ.*культ", l, re.I))
                    for l in current_lessons
                )
                if is_available:
                    slots.setdefault(time_str, []).append(stream_lesson)

def parse_day_date(day_key, ref_date=None):
    if not day_key:
        return None
    if not ref_date:
        ref_date = datetime.date.today()
    m = re.search(r"(\d{1,2})[.](\d{1,2})(?:[.](\d{2,4}))?", day_key)
    if not m:
        return None
    d = int(m.group(1))
    mo = int(m.group(2))
    if m.group(3):
        y = int(m.group(3))
        if y < 100:
            y += 2000
    else:
        y = ref_date.year
        if ref_date.month == 12 and mo == 1:
            y += 1
        elif ref_date.month == 1 and mo == 12:
            y -= 1
    try:
        return datetime.date(y, mo, d)
    except ValueError:
        return None

def is_valid_day_dow(day_key, ref_date=None):
    if not day_key:
        return False
    dt = parse_day_date(day_key, ref_date)
    if not dt:
        return True
    DOW_MAP = {
        'ПОНЕДЕЛЬНИК': 0, 'ВТОРНИК': 1, 'СРЕДА': 2, 'ЧЕТВЕРГ': 3,
        'ПЯТНИЦА': 4, 'СУББОТА': 5, 'ВОСКРЕСЕНЬЕ': 6
    }
    for dow_name, dow_idx in DOW_MAP.items():
        if dow_name in day_key.upper():
            if dt.weekday() != dow_idx:
                return False
    return True

def parse_thursday_kurchatova_block(lessons_list, default_sg1_type):
    # default_sg1_type is 'ПРАКТИКУМ' for morning, 'ПРАКТИКА' for afternoon
    novikova_item = None
    davydova_item = None

    for l in lessons_list:
        clean = re.sub(r'^\d{1,2}[.:]\d{2}\s*[-–]\s*\d{1,2}[.:]\d{2}\s*\n?', '', l).strip()
        clean = re.sub(r'^[12]\s*подгруппа\s*\n?', '', clean).strip()

        # Check if this item is a combined double-lesson cell
        parts = [p.strip() for p in re.split(r'\n(?=ПРАКТИКА|ПРАКТИКУМ)', clean) if p.strip()]
        parts = [p for p in parts if p.startswith('ПРАКТИКА') or p.startswith('ПРАКТИКУМ')]
        if len(parts) >= 2:
            for p in parts:
                if 'Новикова' in p or '716' in p:
                    novikova_item = p
                elif 'Давыдова' in p or '719' in p:
                    davydova_item = p
            if novikova_item and davydova_item:
                break
        else:
            if 'Новикова' in clean or '716' in clean:
                novikova_item = clean
            elif 'Давыдова' in clean or '719' in clean:
                davydova_item = clean

    if not novikova_item or not davydova_item:
        return []

    novikova_clean = re.sub(r'^[12]\s*подгруппа\s*\n?', '', novikova_item).strip()
    davydova_clean = re.sub(r'^[12]\s*подгруппа\s*\n?', '', davydova_item).strip()

    if default_sg1_type == 'ПРАКТИКУМ':
        sg1 = '1 подгруппа\n' + novikova_clean
        sg2 = '2 подгруппа\n' + davydova_clean
    else:
        sg1 = '1 подгруппа\n' + davydova_clean
        sg2 = '2 подгруппа\n' + novikova_clean

    return [sg1, sg2]

def normalize_course1_group9_days(days_dict):
    """
    Course 1 Group 9 travels to ул. Курчатова, 5 (BSU Grove campus) on Mondays and Thursdays.
    They have a dedicated campus bell schedule:
    - Monday (3 pairs):
      Pair 1: 08:30–09:55
      Pair 2: 10:05–11:30
      Pair 3: 11:50–13:15
    - Thursday (4 pairs across 2 subgroups):
      Pair 1: 08:30–09:55
      Pair 2: 10:05–11:30
      Pair 3: 11:50–13:15
      Pair 4: 13:25–14:50
    """
    for day_key in list(days_dict.keys()):
        up = day_key.upper()
        slots = days_dict[day_key]

        if 'ПОНЕДЕЛЬНИК' in up:
            new_slots = {}
            # Slot 1: 08:30–09:55
            s1_items = slots.get('08:30–09:55') or slots.get('08:10–09:30') or []
            s1_clean = []
            for item in s1_items:
                if 'немецк' in item.lower():
                    continue
                c = re.sub(r'^\d{1,2}[.:]\d{2}\s*[-–]\s*\d{1,2}[.:]\d{2}\s*\n?', '', item).strip()
                if 'ул. Курчатова' not in c:
                    c += '\nул. Курчатова, 5'
                if 'ауд.' not in c:
                    c += '\nауд. 718'
                if c not in s1_clean:
                    s1_clean.append(c)
            if s1_clean:
                new_slots['08:30–09:55'] = s1_clean

            # Slot 2: 10:05–11:30
            s2_items = slots.get('10:05–11:30') or slots.get('09:40–11:00') or []
            s2_clean = []
            for item in s2_items:
                if 'немецк' in item.lower():
                    continue
                c = re.sub(r'^\d{1,2}[.:]\d{2}\s*[-–]\s*\d{1,2}[.:]\d{2}\s*\n?', '', item).strip()
                if 'ул. Курчатова' not in c:
                    c += '\nул. Курчатова, 5'
                if 'ауд.' not in c:
                    c += '\nауд. 718'
                if c not in s2_clean:
                    s2_clean.append(c)
            if s2_clean:
                new_slots['10:05–11:30'] = s2_clean

            # Slot 3: 11:50–13:15
            s3_items = slots.get('11:50–13:15') or slots.get('11:20–12:40') or slots.get('12:50–14:10') or []
            s3_clean = []
            for item in s3_items:
                if 'немецк' in item.lower():
                    continue
                c = re.sub(r'^\d{1,2}[.:]\d{2}\s*[-–]\s*\d{1,2}[.:]\d{2}\s*\n?', '', item).strip()
                if 'ул. Курчатова' not in c:
                    c += '\nул. Курчатова, 5'
                if c not in s3_clean:
                    s3_clean.append(c)
            if s3_clean:
                new_slots['11:50–13:15'] = s3_clean

            days_dict[day_key] = new_slots

        elif 'ЧЕТВЕРГ' in up:
            new_slots = {}
            # Morning block (8.30-11.30) -> Pair 1 (08:30–09:55) and Pair 2 (10:05–11:30)
            all_am = []
            for s in ['08:30–09:55', '08:10–09:30', '09:40–11:00', '10:05–11:30']:
                all_am.extend(slots.get(s, []))

            morning_items = parse_thursday_kurchatova_block(all_am, 'ПРАКТИКУМ')
            if morning_items:
                new_slots['08:30–09:55'] = list(morning_items)
                new_slots['10:05–11:30'] = list(morning_items)

            # Afternoon block (11.50-14.50) -> Pair 3 (11:50–13:15) and Pair 4 (13:25–14:50)
            all_pm = []
            for s in ['11:50–13:15', '11:20–12:40', '12:50–14:10', '13:25–14:50', '14:20–15:40']:
                all_pm.extend(slots.get(s, []))

            afternoon_items = parse_thursday_kurchatova_block(all_pm, 'ПРАКТИКА')
            if afternoon_items:
                new_slots['11:50–13:15'] = list(afternoon_items)
                new_slots['13:25–14:50'] = list(afternoon_items)

            days_dict[day_key] = new_slots

def merge_schedule_history(new_data, existing_path="schedule_data.json", max_history_days=7):
    today = datetime.date.today()
    old_data = {}
    if os.path.exists(existing_path):
        try:
            with open(existing_path, "r", encoding="utf-8") as f:
                old_data = json.load(f)
        except Exception as e:
            print(f"Предупреждение: не удалось прочитать {existing_path}: {e}")

    merged_courses = {}
    all_courses = set(old_data.get("courses", {}).keys()) | set(new_data.get("courses", {}).keys())

    def normalize_day_slots(slots_dict):
        normalized = {}
        for slot, lessons in slots_dict.items():
            if not lessons:
                continue
            combined_text = " ".join(lessons)
            custom_time = extract_custom_cell_time(combined_text, slot)
            actual_slot = custom_time if custom_time else slot
            normalized.setdefault(actual_slot, []).extend([l for l in lessons if l not in normalized.get(actual_slot, [])])

        def sort_slot_key(slot_str):
            m = re.match(r"(\d{1,2})[:.](\d{2})", slot_str)
            return int(m.group(1)) * 60 + int(m.group(2)) if m else 9999

        return {k: normalized[k] for k in sorted(normalized.keys(), key=sort_slot_key)}

    for cid in sorted(all_courses):
        merged_courses[cid] = {}
        old_groups = old_data.get("courses", {}).get(cid, {})
        new_groups = new_data.get("courses", {}).get(cid, {})
        all_groups = set(old_groups.keys()) | set(new_groups.keys())

        for g_name in sorted(all_groups):
            merged_days = {}

            # 1. Take historical days that are not older than max_history_days
            for day_key, slots in old_groups.get(g_name, {}).items():
                if not is_valid_day_dow(day_key, today):
                    continue  # Purge invalid typo days like 'СУББОТА 30.10'
                dt = parse_day_date(day_key, today)
                if dt and (today - dt).days > max_history_days:
                    continue  # Prune days older than 7 days
                merged_days[day_key] = normalize_day_slots(slots)

            # 2. Overwrite / append newly parsed days
            for day_key, slots in new_groups.get(g_name, {}).items():
                if not is_valid_day_dow(day_key, today):
                    continue
                dt = parse_day_date(day_key, today)
                if dt and (today - dt).days > max_history_days:
                    continue
                merged_days[day_key] = normalize_day_slots(slots)

            # Clean historical contamination for Group 7 on Monday
            if cid == "1" and g_name == "ГРУППА 7":
                for day_key, slots in merged_days.items():
                    if "ПОНЕДЕЛЬНИК" in day_key.upper():
                        for slot in list(slots.keys()):
                            slots[slot] = [
                                l for l in slots[slot]
                                if not (("ОБЩЕЕ ЗЕМЛЕВЕДЕНИЕ" in l.upper() and slot == "08:10–09:30") or
                                        ("ТОПОГРАФИЯ" in l.upper() and slot == "09:40–11:00") or
                                        ("ОБЩАЯ ЭКОЛОГИЯ" in l.upper() and slot == "11:20–12:40"))
                            ]

            # Clean empty/redundant slots for Group 8 on Monday when 11:50–14:50 is present
            if cid == "1" and g_name == "ГРУППА 8":
                for day_key, slots in merged_days.items():
                    if "ПОНЕДЕЛЬНИК" in day_key.upper():
                        if "11:50–14:50" in slots:
                            if "12:50–14:10" in slots and not slots["12:50–14:10"]:
                                del slots["12:50–14:10"]

            # Subgroup tagging for historical slots:
            # 1. Group 6 Saturday Topography is exclusively for Subgroup 2
            if cid == "1" and g_name == "ГРУППА 6":
                for day_key, slots in merged_days.items():
                    if "СУББОТА" in day_key.upper():
                        for slot in ["09:40–11:00", "11:20–12:40"]:
                            if slot in slots:
                                slots[slot] = [
                                    ("2 подгруппа\n" + l) if not l.startswith("2 подгруппа") and not l.startswith("1 подгруппа") else l
                                    for l in slots[slot]
                                ]

            # 2. When a slot has 2 distinct academic items, slot 0 is 1st subgroup, slot 1 is 2nd subgroup
            for day_key, slots in merged_days.items():
                for slot, lessons in list(slots.items()):
                    # Purge any phantom entries that have no content besides subgroup prefix
                    valid_lessons = []
                    for l in lessons:
                        clean_l = re.sub(r"^[12]\s*подгруппа\s*", "", l, flags=re.I).strip()
                        clean_l = re.sub(r"^\d{1,2}[.:]\d{2}\s*[-–—]\s*\d{1,2}[.:]\d{2}\s*", "", clean_l).strip()
                        clean_l = re.sub(r"^\d{4}\s*[-–—]\s*\d{4}\s*", "", clean_l).strip()
                        if clean_l:
                            valid_lessons.append(l)
                    slots[slot] = valid_lessons
                    lessons = valid_lessons

                    if len(lessons) == 2:
                        l0, l1 = lessons[0], lessons[1]
                        is_lang_pe = any(re.search(r'русск|немец|франц|англ|иностранн|культ', l, re.I) for l in [l0, l1])
                        if not is_lang_pe:
                            if not l0.startswith("1 подгруппа") and not l0.startswith("2 подгруппа"):
                                lessons[0] = "1 подгруппа\n" + l0
                            if not l1.startswith("1 подгруппа") and not l1.startswith("2 подгруппа"):
                                lessons[1] = "2 подгруппа\n" + l1

            # Dedicated bell schedule and pair count for Course 1 Group 9 at ул. Курчатова, 5
            if cid == "1" and g_name == "ГРУППА 9":
                normalize_course1_group9_days(merged_days)

            # Chronological sort of days
            def sort_key(k):
                d = parse_day_date(k, today)
                return d if d else datetime.date.max

            sorted_days = {k: merged_days[k] for k in sorted(merged_days.keys(), key=sort_key)}
            merged_courses[cid][g_name] = sorted_days

    # Merge updated_times (new takes precedence)
    merged_updated_times = dict(old_data.get("updated_times", {}))
    merged_updated_times.update(new_data.get("updated_times", {}))

    return {
        "courses": merged_courses,
        "updated_times": merged_updated_times
    }

def compute_day_hash(slots_dict):
    canonical = json.dumps(slots_dict, sort_keys=True, ensure_ascii=False)
    return hashlib.md5(canonical.encode('utf-8')).hexdigest()[:12]

def finalize_schedule_metadata(final_courses, decanat_times, old_data):
    old_grp_times = old_data.get("group_updated_times", {})
    old_day_times = old_data.get("day_updated_times", {})
    old_day_hashes = old_data.get("day_hashes", {})

    merged_group_updated_times = {}
    merged_day_updated_times = {}
    merged_day_hashes = {}

    for cid, groups in final_courses.items():
        merged_group_updated_times[cid] = {}
        merged_day_updated_times[cid] = {}
        merged_day_hashes[cid] = {}
        course_decanat_time = decanat_times.get(cid, "")

        for g_name, days in groups.items():
            merged_day_updated_times[cid][g_name] = {}
            merged_day_hashes[cid][g_name] = {}
            group_has_changed = False
            prev_grp_time = old_grp_times.get(cid, {}).get(g_name)

            for day_key, slots in days.items():
                cur_h = compute_day_hash(slots)
                merged_day_hashes[cid][g_name][day_key] = cur_h

                old_h = old_day_hashes.get(cid, {}).get(g_name, {}).get(day_key)
                old_dt = old_day_times.get(cid, {}).get(g_name, {}).get(day_key)

                if old_h is None or old_h != cur_h:
                    day_time = course_decanat_time or old_dt or "Обновлено"
                    group_has_changed = True
                else:
                    day_time = old_dt or course_decanat_time or "Обновлено"

                merged_day_updated_times[cid][g_name][day_key] = day_time

            if group_has_changed or not prev_grp_time:
                merged_group_updated_times[cid][g_name] = course_decanat_time or prev_grp_time or "Обновлено"
            else:
                merged_group_updated_times[cid][g_name] = prev_grp_time

    return merged_group_updated_times, merged_day_updated_times, merged_day_hashes

def build_all():
    print("Загрузка файлов деканата...")
    to_convert = []

    for cid, cfg in COURSES_CONFIG.items():
        download_file(cid, cfg)
        if cfg["needs_conversion"]:
            to_convert.append(cfg["raw_file"])

    convert_docs(to_convert)

    # Извлекаем фактическое время последнего обновления файла деканатом
    updated_times = {}
    for cid, cfg in COURSES_CONFIG.items():
        decanat_time = extract_decanat_modtime(cfg["docx_file"])
        if decanat_time:
            updated_times[cid] = decanat_time
            print(f"Курс {cid}: время обновления от деканата = {decanat_time}")

    full_data = {"courses": {}, "updated_times": updated_times}
    for cid, cfg in COURSES_CONFIG.items():
        full_data["courses"][cid] = parse_docx(cfg["docx_file"])

    propagate_stream_languages(full_data["courses"])

    # Load previous schedule_data.json to compare hashes and track per-group/per-day changes
    old_data = {}
    if os.path.exists("schedule_data.json"):
        try:
            with open("schedule_data.json", "r", encoding="utf-8") as f:
                old_data = json.load(f)
        except Exception:
            pass

    final_data = merge_schedule_history(full_data, "schedule_data.json", max_history_days=7)
    if "ГРУППА 9" in final_data["courses"].get("1", {}):
        normalize_course1_group9_days(final_data["courses"]["1"]["ГРУППА 9"])
    propagate_stream_languages(final_data["courses"])
    if "ГРУППА 9" in final_data["courses"].get("1", {}):
        normalize_course1_group9_days(final_data["courses"]["1"]["ГРУППА 9"])

    grp_times, day_times, day_hashes = finalize_schedule_metadata(
        final_data["courses"],
        final_data["updated_times"],
        old_data
    )
    final_data["group_updated_times"] = grp_times
    final_data["day_updated_times"] = day_times
    final_data["day_hashes"] = day_hashes

    with open("schedule_data.json", "w", encoding="utf-8") as f:
        json.dump(final_data, f, ensure_ascii=False, indent=2)

    # Clean up temporary doc and docx files
    for cid, cfg in COURSES_CONFIG.items():
        for fpath in [cfg.get("raw_file"), cfg.get("docx_file")]:
            if fpath and os.path.exists(fpath):
                try:
                    os.remove(fpath)
                except Exception:
                    pass

    print("schedule_data.json успешно собран и очищен от временных файлов!")

if __name__ == "__main__":
    build_all()
